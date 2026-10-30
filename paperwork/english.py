"""English versions of the reviewed worksheet; no language-independent file reuse."""
from pathlib import Path
import json,re,subprocess
from docx import Document
from docx.shared import Pt,RGBColor
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
EN_QUESTION='Using the experimental results and topic notes, explain how bile and lipase in pancreatic juice work together to help digest lipids, and how boiling affects their actions.'
EN_ANSWER='Bile contains no digestive enzymes. It physically emulsifies large oil droplets into small droplets without changing the lipid molecules. The greater total surface area lets lipase contact more lipid. Lipase in pancreatic juice catalyses chemical digestion, producing glycerol and fatty acids in the complete-hydrolysis model used in this course. Lipase is a protein. Boiling denatures it and stops its catalytic action; cooling does not restore it. Bile retains its emulsifying action after boiling in this teaching model.'
TEXT={
'工作紙一：脂質的消化與公平比較':'Worksheet 1: Lipid Digestion and Fair Comparisons',
'工作紙一：脂質的消化與公平比較（參考答案）':'Worksheet 1: Lipid Digestion and Fair Comparisons — Answers',
'姓名：________________　　班別：________　　學號：________':'Name: ________________    Class: ________    No.: ________',
'請先閱讀「課題需知」，再根據提供的情境及模擬結果作答。':'Read the topic notes, then answer using the context and model results.',
'課題需知':'Topic notes',
'01　膽汁：乳化脂質':'01  Bile: emulsifying lipids',
'本教學模型設定消化液 X 是膽汁。膽汁由肝臟製造，儲存於膽囊，不含消化酶。它把大油滴乳化成小油滴，增加同量脂質的總表面積，有助脂肪酶作用。這屬物理消化，脂質分子的結構沒有改變。':'In this model, X is bile. Made in the liver and stored in the gall bladder, bile contains no digestive enzymes. It emulsifies large oil droplets into small droplets, increasing the total surface area for lipase to act on. This is physical digestion: lipid molecules remain unchanged.',
'02　胰液中的脂肪酶：催化化學消化':'02  Lipase in pancreatic juice: chemical digestion',
'本教學模型設定消化液 Y 是胰液，含有脂肪酶。脂肪酶負責催化脂質的化學消化；本課以完全水解的簡化模型，表示甘油三酯分解為甘油與脂肪酸，以助吸收。脂肪酶是蛋白質，高溫可使其變性，失去功能。':'In this model, Y is pancreatic juice containing lipase. Lipase catalyses chemical digestion of lipids. This course uses a simplified complete-hydrolysis model: triglycerides form smaller glycerol and fatty-acid molecules for absorption. Lipase is a protein; high heat can denature it and stop it working.',
'部分 1｜研究情境與公平比較':'Part 1 | Context and fair comparisons',
'研究情境：研究員把消化液 X、Y 加入油水混合物並搖勻，比較各組外觀。':'Researchers add digestive juices X and Y to oil–water mixtures, shake them, and compare their appearance.',
'研究問題：不同消化液組合如何影響油水混合物的外觀？':'Research question: How do combinations of digestive juices affect the appearance of an oil–water mixture?',
'1. 辨識實驗變量':'1. Identify the variables',
'在每列圈出適當的因素；每類可選多於一項。':'Circle the appropriate factors in each row. You may choose more than one.',
'2. 選出公平比較所需的前提（可選多於一項）':'2. Select the assumptions for a fair comparison (choose all that apply)',
'各組起始油量及水量相同，加入液體的總量也相同。':'All sets start with the same amounts of oil and water, and receive the same total volume of added liquid.',
'X＋Y 組可加入較多總液量，以便與其他組比較。':'The X + Y set may receive more total liquid to compare it with the other sets.',
'各組使用相同搖勻方式。':'All sets are shaken in the same way.',
'各組在相同反應時間後觀察。':'All sets are observed after the same reaction time.',
'3. 探究的對照組':'3. Design a control group',
'你認為此實驗需要對照組嗎？如需要，你會如何設計比較裝置？':'Does this experiment need a control group? If so, how would you design the comparison setup?',
'需要。油水混合物加 2 mL 水作對照，與 X、Y 及 X＋Y 組比較；各組起始油水量、加液總量、反應溫度、時間及搖勻方式相同。':'Yes. Add 2 mL of water to an oil–water mixture as a control and compare it with X, Y and X + Y. Keep starting oil/water amounts, total added volume, temperature, time and shaking the same.',
'部分 2.1｜分析四組基本實驗':'Part 2.1 | Analyse the four basic experiments',
'以下是本教學模型提供的模擬結果，並非實際產物檢測。各組依部分 1 的相同條件處理，在相同時間後觀察。':'These are model results, not tests for digestion products. All sets use the same conditions and observation time.',
'4. 消化液 X 有甚麼作用？':'4. What does digestive juice X do?',
'加入消化液 X 後，主要觀察到 （混濁乳狀液 ／ 油層消失）。':'After adding digestive juice X, what is mainly observed? (cloudy emulsion / oil layer disappears).',
'5. 消化液 X 如何有助消化液 Y 作用？':'5. How does digestive juice X help Y act?',
'消化液 X 先把大油滴分散成小油滴，令總表面積 （增加 ／ 減少），有助消化液 Y 作用。':'X disperses large oil droplets into small droplets. Their total surface area (increases / decreases), helping Y act.',
'6. 本實驗能告訴我們甚麼？':'6. What can this experiment tell us?',
'我們只觀察油層是否存在及混濁程度，沒有檢測消化產物。以下哪個說法最恰當？':'We observed only whether an oil layer remained and how cloudy the mixture was, without testing digestion products. Which statement is most appropriate?',
'A. 外觀可提供消化液作用的線索，但不能單獨證明已產生甘油和脂肪酸。':'A. Appearance gives clues about the juices, but alone cannot prove that glycerol and fatty acids have formed.',
'B. 只要油層消失，就已完全證明脂質分解成甘油和脂肪酸。':'B. Disappearance of the oil layer fully proves that lipids have broken down into glycerol and fatty acids.',
'部分 2.2｜加熱對消化液作用的影響':'Part 2.2 | Effects of heating the juices',
'以未煮沸的 X＋Y 組為比較基準，每次只煮沸其中一種消化液。煮沸後先冷卻至 37°C，各組均加入 X、Y 各 1 mL，其餘條件相同。':'Use unboiled X + Y as the baseline. Boil only one juice at a time, then cool it to 37°C. Add 1 mL each of X and Y to every set; keep other conditions the same.',
'7. 延伸結論':'7. Extension conclusion',
'被煮沸後失去作用的消化液是（消化液 X ／ 消化液 Y）。':'The juice that loses its action after boiling is (digestive juice X / digestive juice Y).',
'部分 3｜整合解釋':'Part 3 | Integrated explanation',
'變量類別':'Variable type','可選因素':'Factors to choose from',
'獨立變量\n主動改變的因素':'Independent variable\nFactor changed',
'因變量\n量度或觀察的結果':'Dependent variable\nOutcome measured',
'控制變量\n保持不變的因素':'Controlled variables\nKept the same',
'油和水總體積':'Total oil/water volume','反應溫度':'Reaction temperature','反應時間':'Reaction time','加液總量':'Total added volume','搖勻方式':'Shaking method','混合物外觀':'Mixture appearance','消化液組合':'Combination of digestive juices',
'混濁乳狀液':'cloudy emulsion','增加':'increases','消化液 Y':'digestive juice Y',
}

# Translate the reviewed red terms as well as the surrounding text. Apply these
# only where the Chinese student source has emphasis, never to every occurrence
# across the worksheet (which could inadvertently cue an answer).
KEYWORD_COLOUR='B7312C'
KEYWORDS={
    '膽汁':['bile'],'乳化':['emulsifies'],'小油滴':['small droplets'],
    '總表面積':['total surface area'],'脂肪酶':['lipase'],
    '物理消化':['physical digestion'],'胰液':['pancreatic juice'],
    '化學消化':['chemical digestion'],'甘油三酯':['triglycerides'],
    '甘油':['glycerol'],'脂肪酸':['fatty-acid','fatty acids'],
    '蛋白質':['protein'],'變性':['denature'],'對照組':['control group'],
    '獨立變量':['Independent variable'],'因變量':['Dependent variable'],
    '控制變量':['Controlled variables'],
}

def red_terms(paragraph):
    return [r.text for r in paragraph.runs
            if r.text.strip() and str(r.font.color.rgb)==KEYWORD_COLOUR]

def translate(source,dictionary):
    if not source.strip():return source
    if source.startswith('8. '):return '8. '+EN_QUESTION
    if source.startswith('膽汁不含消化酶，透過物理消化'):return EN_ANSWER
    prefix=''
    if source.startswith(('□ ','☑ ','✓ ')):prefix,source=source[:2],source[2:]
    if source in ['油和水總體積','反應溫度','反應時間','加液總量','搖勻方式','混合物外觀','消化液組合'] and source in dictionary:return prefix+dictionary[source]
    if source in TEXT:return prefix+TEXT[source]
    if source in dictionary:return prefix+dictionary[source]
    if source.startswith('（') and source.endswith('）') and source[1:-1] in dictionary:return '('+dictionary[source[1:-1]]+')'
    if ' ／ ' in source or '\n' in source:
        pieces=re.split(r'( ／ |\n)',source)
        return ''.join(' / ' if x==' ／ ' else '\n' if x=='\n' else translate(x,dictionary) for x in pieces)
    raise ValueError('Untranslated worksheet text: '+source)

def create(root,diagram_dir):
    command="const fs=require('fs'),vm=require('vm');const c={window:{}};vm.runInNewContext(fs.readFileSync('language-dictionary.js','utf8'),c);process.stdout.write(JSON.stringify(c.window.VL1_TRANSLATIONS.texts))"
    dictionary=json.loads(subprocess.check_output(['node','-e',command],cwd=ROOT,text=True))
    # The answer generator may rebuild a choice paragraph to mark its answer.
    # Use the reviewed student text for the shared conceptual emphasis too.
    emphasis={p.text:red_terms(p) for p in Document(root/'WS1_student.docx').paragraphs}
    for stem in ['student','answers']:
        doc=Document(root/f'WS1_{stem}.docx')
        for style_name in ['Normal','Heading 1','Heading 2','Heading 3']:
            style=doc.styles[style_name];style.font.name='Arial'
            fonts=style._element.get_or_add_rPr().rFonts
            fonts.set(qn('w:ascii'),'Arial');fonts.set(qn('w:hAnsi'),'Arial')
        paragraphs=list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:paragraphs.extend(cell.paragraphs)
        for paragraph in paragraphs:
            original=paragraph.text
            if not original.strip():continue
            english=translate(original,dictionary)
            runs=[r for r in paragraph.runs if r.text.strip()]
            answer=bool(runs) and all(r.font.color.rgb and str(r.font.color.rgb)=='145DA0' for r in runs)
            chosen=[translate(r.text,dictionary) for r in runs if r.underline and r.font.color.rgb and str(r.font.color.rgb)=='145DA0']
            terms=emphasis.get(original,red_terms(paragraph))
            keywords=[word for term in terms for word in KEYWORDS[term]]
            keyword_pattern=r'\b(?:'+'|'.join(re.escape(x) for x in keywords)+r')\b' if keywords else None
            hint=original.startswith('（')
            heading=paragraph.style.name.startswith('Heading')
            paragraph.clear()
            pieces=[re.escape(x) for x in sorted(chosen,key=len,reverse=True)]
            if keyword_pattern:pieces.append(keyword_pattern)
            pattern='('+'|'.join(pieces)+')' if pieces else None
            for part in re.split(pattern,english,flags=re.IGNORECASE) if pattern else [english]:
                if not part:continue
                run=paragraph.add_run(part)
                if not heading:run.font.size=Pt(11 if hint else 12)
                if answer or part in chosen:
                    run.font.color.rgb=RGBColor.from_string('145DA0');run.bold=True;run.underline=part in chosen
                elif keyword_pattern and re.fullmatch(keyword_pattern,part,re.IGNORECASE):
                    run.font.color.rgb=RGBColor.from_string(KEYWORD_COLOUR);run.bold=True
                elif hint:run.font.color.rgb=RGBColor.from_string('658087')
            if english.startswith('Name:'):paragraph.paragraph_format.line_spacing=1.2
            if original.startswith('你認為此實驗需要對照組嗎？') or original.startswith('（提示：對照組'):
                paragraph.paragraph_format.keep_with_next=True
        for section in doc.sections:
            h=section.header.paragraphs[0];h.text='Biology | Lipid digestion and fair comparisons'+(' | ANSWERS' if stem=='answers' else '')
            for run in h.runs:run.font.size=Pt(11)
            for paragraph in section.footer.paragraphs:
                for run in paragraph.runs:
                    if '工作紙一' in run.text:run.text=run.text.replace('工作紙一','Worksheet 1')
                    run.font.size=Pt(11)
        names=['01_emulsification','02_hydrolysis','03_scenario','04_basic_results','05_heating_results']
        for shape,name in zip(doc.inline_shapes,names):
            rid=shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
            doc.part.related_parts[rid]._blob=(diagram_dir/(name+'.png')).read_bytes()
        doc.core_properties.title='Worksheet 1: Lipid Digestion and Fair Comparisons'+(' — Answers' if stem=='answers' else '')
        doc.save(root/f'WS1_{stem}_en.docx')
