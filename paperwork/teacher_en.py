"""English counterpart of the current teacher guide and its versioned rubric."""
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from english import EN_QUESTION
QUESTIONS=[
('Identify variables','Design: IV / DV / CV',4,'IV: combination of digestive juices. DV: mixture appearance. CV: total oil/water volume, temperature, reaction time, total added volume and shaking method.','IV 1; DV 1; CV 2. Each correct CV earns 0.4. Selecting a non-CV makes the CV component 0.'),
('Assumptions for fairness','Design: assumptions',1,'Choose statements 1, 3 and 4, not extra liquid for X + Y.','All three correct with no wrong selection: 1. Other attempted combinations: 0.'),
('Control-group design','Design: control plan',2,'Add 2 mL of water to an oil–water mixture as a control; compare with X, Y and X + Y under equal conditions.','Suitable water control and comparison purpose: 1. Fair starting amounts, added volume, temperature, time and shaking: 1. Accept other sound wording.'),
('Action of X','Analysis: q1',1,'Cloudy emulsion.','Correct: 1; attempted wrong answer: 0. Do not add a physical/chemical classification or written comparison.'),
('How X helps Y','Analysis: q3',1,'Increases.','Correct: 1; attempted wrong answer: 0. No additional molecular-structure explanation is required.'),
('Limits of appearance','Analysis: limitations',1,'A: appearance gives clues but cannot alone prove that glycerol and fatty acids formed.','A: 1; attempted wrong answer: 0.'),
('Effect of boiling','Extension: q2',1,'Digestive juice Y.','Y: 1; attempted wrong answer: 0. No extra control or cooling explanation is required here.'),
('Integrated concepts','Learning reflection: same concepts, with topic notes named in the worksheet',8,'See the complete answer version.','Use the shared eight-point concept rubric. Do not add observation comparisons, appearance limits or hypothesis revision.')]
POINTS=[
'Bile physically emulsifies large oil droplets into small droplets.',
'Emulsification does not change the structure of lipid molecules.',
'Smaller droplets increase the total surface area for lipase to contact the same amount of lipid.',
'Lipase in pancreatic juice catalyses chemical digestion of lipids.',
'The course complete-hydrolysis model forms glycerol and fatty acids.',
'Lipase is a protein.',
'Boiling denatures lipase and stops its catalytic function. Failure to recover after cooling is acceptable supporting detail.',
'Bile retains its emulsifying function after boiling in this teaching model.']
SPS=[
('Observation','Initial observation: teacher 0–2. Six-set appearances: automatic 0–2.','An accurate specific visible description earns 2; partial description 1; no valid observation 0. After all six sets are answered, appearances earn correct sets / 6 × 2; missing sets are unavailable, not invented zeros.'),
('Classification','IV 1 + DV 1 + CV 2, automatic.','IV is juice combination, DV is appearance. Five CVs earn 0.4 each; selecting a non-CV makes the CV component 0.'),
('Planning','Original prediction/reason: teacher 0–2; assumptions: automatic 0–1; control: teacher 0–1.','A testable prediction with a sound reason earns 2; partial evidence 1. A reasonable prediction is not wrong just because results differ. The control must have a comparison purpose and a sound setup.'),
('Carrying out investigations','Setup quality: teacher 0–2; fair conditions and arrangements: teacher 0–2.','Judge clear labels/volumes and a workable plan, including fair treatment and cooling. Partial evidence earns 1 in each component. Clicks, speed, completion and time are not direct ability scores.'),
('Inference','q1, q3, limitations and q2: automatic, 1 each.','Use the answers for the record version. An attempted wrong answer scores 0; missing evidence remains unavailable.'),
('Communication','Teacher 0–4, from the reasons, control/setup explanations, data and reflection actually requested.','4: clear organisation and causal explanation; 3: minor missing links; 2: understandable but incomplete; 1: one relevant statement; 0: no usable communication. Do not require extra comparisons or limits in the current final question or double-count concept accuracy.')]

def create(path):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
 s.top_margin=s.bottom_margin=Cm(1.5);s.left_margin=s.right_margin=Cm(1.5);s.header_distance=s.footer_distance=Cm(.7)
 for name in ['Normal','Heading 1','Heading 2']:
  style=d.styles[name];style.font.name='Noto Sans';style.font.size=Pt(12);style.paragraph_format.line_spacing=1.25;style.paragraph_format.space_after=Pt(6)
 d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 2'].font.size=Pt(13)
 for name in ['Heading 1','Heading 2']:d.styles[name].font.color.rgb=RGBColor.from_string('245D60')
 h=s.header.paragraphs[0];h.text='VL1 | D · Teacher guide, mapping and marking | TEACHER USE';h.runs[0].font.size=Pt(11)
 f=s.footer.paragraphs[0];f.alignment=2;f.add_run('D | ').font.size=Pt(11);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 def title(t,first=False):
  p=d.add_paragraph(t,'Heading 1');p.paragraph_format.page_break_before=not first
 def sub(t):d.add_paragraph(t,'Heading 2')
 def p(t):d.add_paragraph(t)
 def table(headers,rows):
  t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
  for c,v in zip(t.rows[0].cells,headers):c.text=v
  for values in rows:
   for c,v in zip(t.add_row().cells,values):c.text=str(v)
  for i,row in enumerate(t.rows):
   row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
   for c in row.cells:
    for para in c.paragraphs:
     para.paragraph_format.line_spacing=1.15;para.paragraph_format.space_after=Pt(3)
     for run in para.runs:run.font.size=Pt(11);run.bold=i==0
 title('D · Teacher guide, mapping and marking',True)
 p('For the reviewed eight-question worksheet (student and answer copies, four pages each) and the current VL. This guide is for teachers, not an additional student task.')
 sub('Materials and distribution')
 p('Give students the full blank worksheet in PDF or Word. Retain the answer version and this guide. Save first answers before showing the answer copy; students correct in another colour without erasing the original.')
 p('The panel has five choices: student PDF/Word, complete answer PDF/Word, and this D guide PDF. Switching the VL language changes all five to the corresponding language. Old A–C parts and the old full pack are retired.')
 sub('Common timing and support')
 p('Choose the same total T of 15–20 minutes for both groups, including reading, answers and checking. For 20 minutes, an example is 5 / 13 / 2; for 15, 4 / 9 / 2. Trial with the class and record actual start/end times; this is a schedule, not proof of feasibility for every student.')
 p('The comparison group learns from the supplied notes, diagrams and results without operating the VL or receiving a concept pre-lesson. Clarify wording/procedure without supplying first-answer solutions; record extra support consistently.')
 sub('Science and controls')
 p('X is bile; Y is pancreatic juice containing lipase. Lipase is an enzyme/protein, not a digestive juice. Emulsification is physical; the hydrolysis diagram is a simplified complete-hydrolysis model. Appearance alone cannot prove products or rates.')
 p('Worksheet hints match the VL. Explain controls for the question: water for the main comparison; unboiled X + Y for heating. A control does not universally exclude X or Y. Do not add student follow-up questions.')
 title('Mapping of the eight worksheet questions')
 p('The worksheet does not assess all VL inputs or operations. Do not add questions or assign zeros just to create matching assessment opportunities.')
 table(['Worksheet','VL counterpart','WS marks'],[(f'{i}. {q[0]}',q[1],q[2]) for i,q in enumerate(QUESTIONS,1)])
 sub('Mapping is not score equivalence')
 p('WS questions 1–7 total 11; question 8 adds 8, for 19. The current VL totals 32 (SPS 24 + reflection 8) and includes observation, prediction and practical evidence absent from the worksheet. Do not convert the totals proportionally; use the study’s separate common assessments for research outcomes.')
 p('Question 8 and the VL reflection assess the same concepts; the worksheet says “topic notes”, while the VL says “key learning points”. Both use the same eight concept points.')
 title('Worksheet marking: questions 1–7')
 for i,q in enumerate(QUESTIONS[:7],1):sub(f'{i}. {q[0]} ({q[2]} marks)');p(q[3]+' '+q[4])
 p('Mark missing responses unavailable/pending according to the study’s missing-data rule. Only attempted wrong answers receive zero. Accept scientifically sound alternative wording in open responses.')
 title('Integrated explanation: 8 marks')
 p('Worksheet prompt: '+EN_QUESTION)
 p('The VL wording uses “key learning points” in place of “topic notes”.')
 sub('One mark for each correct concept')
 for i,point in enumerate(POINTS,1):p(f'{i}. {point}')
 sub('Partial credit and omissions')
 p('Award only concepts expressed correctly, accepting equivalent wording rather than counting keywords. Protein identity is separate from the causal link between denaturation and loss of catalytic action. Cooling is supporting detail, not an extra required question.')
 p('Do not additionally require observation comparisons, appearance limitations or revision of the first hypothesis. Question 6 already assesses appearance limits. Missing answers remain blank; attempted answers with no correct concepts score 0. This post-learning item alone does not prove knowledge gain.')
 title('VL: six SPS areas, 24 marks')
 for name,evidence,criteria in SPS:sub(name+' (0–4)');p(evidence+' '+criteria)
 p('This is a draft activity-feedback rubric, not a validated SPSAI/CKT. Calibrate opportunities, difficulty and rater agreement before comparisons across modules, groups or versions.')
 title('Excel fields, versions and saving')
 sub('Use the rubric for each record')
 p('The teacher-scoring sheet includes a current/previous reflection score, rubric version and activity maximum. Current integrated-concept and earlier integrated-evidence tasks use 0–8; the earlier denaturation-only task uses 0–4. Fill the six applicable SPS manual cells and this reflection cell; the old four reflection cells are not applicable. Original legacy records retain their four 0–2 reflection cells instead.')
 p('The sheet names remain those of the shared Chinese class export: student answers; six-set observations; teacher scoring; rubric; events; original/submitted snapshots; setup images. The material language switch does not translate or rewrite stored student responses or the shared research workbook.')
 sub('Blank, zero and invalid values')
 p('Blank means pending; 0 is a marked zero. Applicable manual fields accept integers within bounds. Formulas also reject pasted negative, excessive, fractional or text values. An overall total appears only after reflection submission and completion of all applicable marks. Inactive cells are not mandatory.')
 sub('Records and colours')
 p('Original, first and final evidence and historical records remain intact. Green means full/correct, orange partial and red zero/wrong. Reflection colours use that row’s maximum of 4 or 8. Student PDFs contain no ability marks.')
 sub('Version limits and saving')
 p('Current integrated-concept records use integrated-concept-explanation (32 total). Earlier integrated-explanation-evidence records retain 32, enzyme-denaturation retains 28, and original untagged legacy records retain 32. Preserve their original questions and criteria; equal maxima do not make the scales equivalent.')
 p('Save the original export and mark a copy. Teacher scores exist only in the separately saved Excel; re-export does not retrieve previously entered marks. Existing downloaded files do not update automatically. WS 19 marks are separate paper feedback and do not fill VL SPS fields for work never asked on paper.')
 d.core_properties.title='D · Teacher guide, mapping and marking';d.save(path)
