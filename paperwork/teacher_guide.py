"""Teacher-only guide derived from the current eight questions and versioned rubric."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
HERE=Path(__file__).resolve().parent

def create(path):
    rules=json.loads((HERE/'rubric.json').read_text())
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=s.bottom_margin=Cm(1.5);s.left_margin=s.right_margin=Cm(1.5)
    s.header_distance=s.footer_distance=Cm(.7)
    for name in ['Normal','Heading 1','Heading 2']:
        style=d.styles[name];style.font.name='Noto Sans CJK TC';style.font.size=Pt(12)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK TC')
        style.paragraph_format.line_spacing=1.3;style.paragraph_format.space_after=Pt(6)
    d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 2'].font.size=Pt(13)
    for name in ['Heading 1','Heading 2']:d.styles[name].font.color.rgb=RGBColor.from_string('245D60')
    head=s.header.paragraphs[0];head.text='VL1｜D · 教師指引、配對及評分｜教師用';head.runs[0].font.size=Pt(11)
    foot=s.footer.paragraphs[0];foot.alignment=2;foot.add_run('D　｜　').font.size=Pt(11)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    def title(t,new=True):
        p=d.add_paragraph(t,'Heading 1');p.paragraph_format.page_break_before=new
    def h(t):d.add_paragraph(t,'Heading 2')
    def p(t):d.add_paragraph(t)
    def table(headers,rows,widths):
        t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
        for c,w in zip(t.columns,widths):c.width=Cm(w)
        for cell,text in zip(t.rows[0].cells,headers):cell.text=text
        for values in rows:
            for cell,text in zip(t.add_row().cells,values):cell.text=str(text)
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for j,c in enumerate(row.cells):
                c.width=Cm(widths[j])
                for para in c.paragraphs:
                    para.paragraph_format.line_spacing=1.2;para.paragraph_format.space_after=Pt(3)
                    for run in para.runs:run.font.size=Pt(11);run.bold=i==0
        return t
    title('D · 教師指引、配對及評分',False)
    p('適用：教師核訂的八題完整工作紙（學生版／答案版各四頁）及現版 VL。此文件供教師使用，不另派作學生題目。')
    h('材料與派發')
    p('學生使用完整工作紙 PDF 或 Word；教師保留連答案版本及本指引。開始先派學生版；保存首次答案後，才使用答案版核對。學生用另一顏色修正，保留原答。')
    p('網站只提供五項下載：完整工作紙 PDF、Word；完整工作紙連參考答案 PDF、Word；本 D 指引 PDF。舊 A、B、C 及完整教師套裝不再提供。')
    h('共同時間與支援')
    p('兩組總時間均為 15–20 分鐘，由研究者選定同一 T，包含閱讀、作答及核對。例：T＝20 分鐘時閱讀 5、作答 13、核對 2；T＝15 時可用 4、9、2。正式使用前按班級閱讀速度試行，記錄實際起止時間。')
    p('比較組毋須操作 VL 或接受課前概念講授，依已提供的概念、圖示及模擬結果作答。教師可澄清文字／程序，不在首次作答前給答案；額外支援依兩組共同規則記錄。')
    h('科學說明與課堂界限')
    p('X 是膽汁；Y 是含脂肪酶的胰液。脂肪酶是酶／蛋白質，不是消化液。膽汁乳化屬物理作用；水解圖採完全水解簡化模型。外觀不能單獨證明產物或消化速度。')
    p('工作紙保留與 VL 相同的提示。教師須按研究問題解釋對照：主比較用水組；熱處理以未煮沸 X＋Y 為基準。對照組並非一概不含 X／Y。不新增學生追問。')
    title('八題逐題配對與作答證據')
    p('下表只列現有八題。VL 額外的初步觀察、原始預測及裝置操作不在本 WS 評分範圍；不得為補分增加未問的問題。')
    table(['WS 題目','VL 對應','WS 分'],[(f"{q['number']}. {q['title']}",q['vl'],q['max']) for q in rules['worksheet']],[6.0,10.0,2.0])
    h('配對不等於總分相等')
    p('WS 練習共 19 分；現版 VL 活動共 32 分，包含紙本沒有的原始預測、觀察及實作證據。兩者不能直接互換或按比例換算。共同成效應另用研究既定的共同測驗評估。')
    p('WS 第 8 題及 VL 學習反思是同一概念任務（VL 指「學習重點」、WS 指「課題需知」），使用相同八點／8 分規準。')
    title('WS 批改：第 1–7 題')
    for q in rules['worksheet'][:7]:
        h(f"{q['number']}. {q['title']}（{q['max']} 分）")
        p(q['answer']+' '+q['marking'])
    p('缺答標示「未提供／待評」，依研究缺失資料規則處理，不捏造作答。已作答但錯誤才按該項給 0。開放題接受其他科學上合理表達。')
    title('共同反思／整合解釋：8 分')
    p('VL 題目：'+rules['question'])
    p('WS 題幹同上，將「學習重點」改為「課題需知」。')
    h('逐項給分，每項 1 分')
    for i,point in enumerate(rules['reflectionPoints'],1):p(f'{i}. {point}')
    p('參考解釋見完整答案版；本頁規準與現版 Excel 的綜合解釋 0–8 欄一致。')
    h('部分得分及缺漏')
    p('正確解釋幾項便給幾分，可用同義措辭，不靠關鍵字數量。高溫一項需連結變性及失去催化功能；蛋白質身分另計一項。僅列關鍵字而無正確表述，不自動得分。')
    p('本題涵蓋膽汁、乳化與分子結構、表面積、脂肪酶及產物、煮沸影響。不另要求引用組別觀察、外觀限制或修訂假說；外觀限制仍由第 6 題評分。未答保留空白；已答無正確概念為 0。此單題不能單獨證明知識新增。')
    title('VL：六項 SPS（共 24 分）')
    for row in rules['sps']:
        h(row['name']+'（0–4）')
        p(row['evidence']+' '+row['criteria'])
    p('此為活動回饋 rubric 初稿，並非已驗證 SPSAI／CKT；跨 VL、組別及版本比較須先校準題目機會、難度與評分者一致性。')
    title('Excel 欄位、公式與保存')
    h('按每列題目版本評分')
    p('「教師評分」提供綜合概念／前題反思欄、反思評分版本及活動滿分。現版綜合題該欄限 0–8；酶變性前題限 0–4。兩者只填六個 SPS 人工欄及此欄，舊四欄顯示「不適用」。原始假說舊題則保留舊四欄，新欄不適用。')
    p('自動題按答案內容計算；人工欄不自動判斷學生文字。現版反思總分只引用現版 0–8 欄，不要求填舊欄。舊版仍引用四個 0–2 欄。')
    h('空白、零分及無效數值')
    p('空白是待評，0 是已評零分。適用人工欄有整數範圍驗證，公式亦檢查貼上的負數、超上限、非整數及文字；無效值不得產生有效總分。學生提交反思且適用評分完整才顯示總分。')
    h('紀錄及色彩')
    p('原有七張工作表保留：學生探究答案、六組觀察紀錄、教師評分、評分準則、操作事件紀錄、原始與遞交快照、裝置設計圖。原始答案、首次與最後觀察及歷史紀錄不改寫。')
    p('綠字表示滿分／正確，橙字表示部分分，紅字表示零分／錯誤；反思顏色按該列 4 或 8 分上限判斷。不適用欄不列為必填評分。學生 PDF 不顯示能力分數。')
    h('教師操作及版本限制')
    p('先保存原始匯出檔，再於副本填分並另存。人工分數只在教師另存的 Excel；重新下載不會讀回已填分。既有已下載 Excel 不會自動更新，須從更新後 VL 重新匯出。')
    p('WS 19 分是獨立紙本回饋，不會寫入 VL 學生紀錄，也不會填入學生沒有作答機會的 SPS 欄。現版綜合概念題 32、證據綜合前題 32、酶變性前題 28、原始假說舊題 32；題目與證據不同，不因滿分相同而當作同一尺度，不按新規準重評舊答。')
    d.core_properties.title='D · 教師指引、配對及評分';d.save(path)
