"""Validate the actual browser export. Run after smoke.cjs; needs openpyxl."""
import re
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

path = '/tmp/vl1-records.xlsx'
with ZipFile(path) as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        if name.endswith(('.xml', '.rels')):
            ET.fromstring(archive.read(name))
    assert any(n.startswith('xl/media/') for n in archive.namelist())
w = load_workbook(path)
assert w.sheetnames == ['學生探究答案', '六組觀察紀錄', '教師評分', '評分準則', '操作事件紀錄', '原始與遞交快照', '裝置設計圖']
answers, scores, observations = w['學生探究答案'], w['教師評分'], w['六組觀察紀錄']
assert answers['B2'].value == '陳小明'
heads={c.value:c.column for c in answers[1]}
assert answers.cell(2,heads['主探究結論｜Y 組的外觀']).font.color.rgb=='FF00834A'
assert answers.cell(2,heads['實驗版本']).value=='staged-4'
assert scores['AG2'].value==32
assert len(w['裝置設計圖']._images) == 1
assert answers['N2'].font.color.rgb == 'FF00834A'  # Independent variable, correct.
assert answers['Q2'].font.color.rgb == 'FFC03030'  # Deliberately incorrect premise selection.
assert answers['I2'].font.color.rgb == 'FF173E34'  # Open observation remains unmarked.
assert len({answers[f'{c}2'].fill.fgColor.rgb for c in ['I', 'N', 'Q', 'T', 'X', 'Y']}) == 6
assert observations['D2'].font.color.rgb == 'FFC03030'  # First X answer was incorrect.
assert observations['E2'].font.color.rgb == 'FF00834A'  # Final X answer was correct.
assert len(scores.data_validations.dataValidation) == 31
assert len(answers.conditional_formatting) == 4 * 6
assert scores['N2'].value in ('',None)
assert 'N2' in [str(v.sqref) for v in scores.data_validations.dataValidation]
assert w.calculation.fullCalcOnLoad and w.calculation.forceFullCalc
assert w['評分準則'].max_row >= 19
assert w['操作事件紀錄'].max_row > 80
assert w['原始與遞交快照'].max_row >= 9
assert all(len(str(c.value or '')) <= 32767 for sheet in w for row in sheet for c in row)
manual = [str(v.sqref).split(':')[0] for v in scores.data_validations.dataValidation]
assert all(scores[c].value in ('', None) for c in manual)
assert all(v.errorStyle in (None, 'stop') and v.allowBlank and v.showErrorMessage for v in scores.data_validations.dataValidation)

# Small parser evaluates the actual formulas, including lazy IF / IFERROR.
# It is not a replacement for Excel; only the exported subset is interpreted.
def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)

TOKEN = re.compile(r'\s*("(?:[^"]|"")*"|[A-Z]+[0-9]+|[A-Z][A-Z0-9_]*|\d+(?:\.\d+)?|>=|<=|<>|[=><(),])')

def parse(formula):
    tokens = TOKEN.findall(formula)
    assert ''.join(tokens) == re.sub(r'\s+(?=(?:[^"]*"[^"]*")*[^"]*$)', '', formula)
    pos = 0
    def atom():
        nonlocal pos
        token = tokens[pos]; pos += 1
        if token.startswith('"'): return ('value', token[1:-1].replace('""', '"'))
        if re.fullmatch(r'\d+(?:\.\d+)?', token): return ('value', float(token))
        if re.fullmatch(r'[A-Z]+[0-9]+', token): return ('ref', token)
        if token in ('TRUE', 'FALSE'): return ('value', token == 'TRUE')
        assert tokens[pos] == '('; pos += 1
        args = []
        if tokens[pos] != ')':
            args.append(expr())
            while tokens[pos] == ',':
                pos += 1; args.append(expr())
        assert tokens[pos] == ')'; pos += 1
        return ('call', token, args)
    def expr():
        nonlocal pos
        left = atom()
        if pos < len(tokens) and tokens[pos] in ('=', '>=', '<=', '>', '<', '<>'):
            op = tokens[pos]; pos += 1
            return ('op', op, left, atom())
        return left
    root = expr(); assert pos == len(tokens)
    return root

def ev(cell, seen=frozenset()):
    assert cell not in seen, 'circular reference'
    value = scores[cell].value
    if not isinstance(value, str) or not value.startswith('='): return '' if value is None else value
    return evaluate(parse(value[1:]), seen | {cell})

def evaluate(node, seen):
    kind = node[0]
    if kind == 'value': return node[1]
    if kind == 'ref': return ev(node[1], seen)
    if kind == 'op':
        op, a, b = node[1], evaluate(node[2], seen), evaluate(node[3], seen)
        if op == '=': return a == b
        if op == '<>': return a != b
        if number(a) != number(b): return False
        return {'>': lambda: a > b, '<': lambda: a < b, '>=': lambda: a >= b, '<=': lambda: a <= b}[op]()
    _, fn, args = node
    if fn == 'IF': return evaluate(args[1] if evaluate(args[0], seen) else args[2], seen)
    if fn == 'IFERROR':
        try: return evaluate(args[0], seen)
        except (TypeError, ValueError, ZeroDivisionError): return evaluate(args[1], seen)
    values = [evaluate(arg, seen) for arg in args]
    funcs = {'AND': lambda *xs: all(xs), 'ISNUMBER': number,
             'SUM': lambda *xs: sum(x for x in xs if number(x)),
             'COUNT': lambda *xs: sum(number(x) for x in xs),
             'ROUND': lambda x, n: round(x, int(n)), 'MOD': lambda x, y: x % y}
    return funcs[fn](*values)

assert ev('AC2') == '待評／未完成'
assert ev('W2') == '待評'  # SPS total.
assert ev('AB2') == '待評'  # Knowledge.
assert ev('G2') == '待評'  # Observation requires manual evidence.
assert ev('K2') == 4
assert ev('M2') == 0  # Incorrect assumption set.
assert ev('S2') == 4
for validation in scores.data_validations.dataValidation:
    address=str(validation.sqref).split(':')[0]
    if re.search(r'(?<![0-9])2$',address): scores[address] = int(validation.formula2)
assert ev('AB2') == 8
assert ev('W2') == 23  # One point lost on premises.
assert ev('AC2') == 31
assert ev('AD2') == '評分完成'
scores['AE2'] = 0
assert ev('AB2') == 0
scores['AE2'] = None
assert ev('AB2') == '待評' and ev('AC2') == '待評／未完成'
for invalid in (9, -1, 1.5, '8', '待評'):
    scores['AE2'] = invalid
    assert ev('AB2') == '待評' and ev('AC2') == '待評／未完成'
scores['AE2'] = 8; scores['D2'] = '待提交反思'
assert ev('AC2') == '待評／未完成'
scores['D2'] = '已完成'

# Conditional answer colours use manual scores while preserving category fills.
def colour_matches(target):
    rules = next(rules for cf, rules in answers.conditional_formatting._cf_rules.items() if str(cf.sqref) == target)
    matches = []
    for rule in rules:
        formula = rule.formula[0]
        # Substitute INDIRECT with the evaluated referenced score, as a literal.
        def literal(match):
            v = ev(match[1])
            return str(v) if number(v) else '"' + str(v).replace('"', '""') + '"'
        formula = re.sub(r'INDIRECT\("\'教師評分\'!([A-Z]+[0-9]+)"\)', literal, formula)
        if evaluate(parse(formula), frozenset()): matches.append(rule.dxfId)
    return matches
scores['E2'] = None; assert colour_matches('I2') == []
scores['E2'] = 2; assert colour_matches('I2') == [0]
scores['E2'] = 0; assert colour_matches('I2') == [1]
scores['E2'] = 1; assert colour_matches('I2') == [2]
scores['AE2'] = 0
assert colour_matches('Y2') == [1]
scores['AE2'] = None; assert colour_matches('Y2') == []
# Legacy missing new answers must not manufacture scores or original responses.
assert answers['J5'].value == '未提供（舊版未保存原始假說）'
assert scores['J5'].value == '未提供'
assert scores['M5'].value == '未提供'
assert scores['S5'].value == '未提供'
assert ev('AC5') == '待評／未完成'
print('PASS: seven XLSX sheets, XML integrity, styles, first/final observation colours, images, version-specific validated manual fields, six SPS totals, pending vs zero, bounds, knowledge, overall gate, legacy missing values and conditional colours.')

# edge_cases.cjs exports a payload larger than one Excel cell, plus formula-like text.
from pathlib import Path
large_path = Path('/tmp/vl1-large.xlsx')
if large_path.exists():
    large = load_workbook(large_path)
    with ZipFile(large_path) as archive:
        for name in archive.namelist():
            if name.endswith(('.xml', '.rels')): ET.fromstring(archive.read(name))
    event_rows = list(large['操作事件紀錄'].iter_rows(min_row=2, values_only=True))
    import json
    payload = Path('/tmp/vl1-large-payload.txt').read_text()
    event = json.loads(''.join(row[7] for row in sorted(event_rows, key=lambda row: row[6])))
    assert event['value'] == payload
    snapshot_rows = [row for row in large['原始與遞交快照'].iter_rows(min_row=2, values_only=True) if row[2] == '第一次實驗前固定快照']
    snapshot = json.loads(''.join(row[5] for row in sorted(snapshot_rows, key=lambda row: row[4])))
    assert snapshot['form']['reason'] == payload
    assert large['學生探究答案']['C2'].data_type == 's'
    assert large['學生探究答案']['C2'].value == '=1+1'
    assert all(len(str(c.value or '').encode('utf-16-le')) // 2 <= 32767 for sheet in large for row in sheet for c in row)
    print('PASS: large Unicode event/snapshot reconstructed without loss, XML valid, formula-like user text kept as text.')

# Evaluate actual Excel formulas for the new comprehensive question, the interim
# denaturation-only question, and original historical reflection in one workbook.
version_path=Path('/tmp/vl1-rubric-versions.xlsx')
assert version_path.exists(), 'Run tests/scoring_versions.cjs first'
mixed=load_workbook(version_path);scores=mixed['教師評分']
assert len(scores.data_validations.dataValidation)==57
for row,maximum,reflection_max in [(2,32,8),(3,32,8),(4,31,8),(5,31,8),(6,32,8),(7,32,8),(8,28,4),(9,32,8)]:
    assert scores[f'AG{row}'].value==maximum
    assert ev(f'AC{row}')=='待評／未完成'
    applicable=[v for v in scores.data_validations.dataValidation if str(v.sqref).endswith(str(row))]
    assert len(applicable)==(6 if row in [4,5] else 10 if row==9 else 7)
    for v in applicable:scores[str(v.sqref)]=int(v.formula2)
    assert ev(f'AB{row}')==reflection_max and ev(f'AC{row}')==maximum
    assert ev(f'AD{row}')=='評分完成'
    answer_cell=f'X{row}' if row==9 else f'AE{row}'
    limit=2 if row==9 else reflection_max
    scores[answer_cell]=0
    assert ev(f'AB{row}')==(6 if row==9 else 0)
    for invalid in [None,limit+1,-1,0.5,'2']:
        scores[answer_cell]=invalid
        assert ev(f'AB{row}')=='待評' and ev(f'AC{row}')=='待評／未完成'
    scores[answer_cell]=limit
    assert ev(f'AC{row}')==maximum
    scores[f'D{row}']='待提交反思'
    assert ev(f'AC{row}')=='待評／未完成'
    scores[f'D{row}']='已完成'
print('PASS: actual mixed-version formulas: 31/32/32/28/32 maxima, 8/8/8/4/8 reflection, inactive fields, blank/zero/invalid values, submission gates and legacy preservation.')
