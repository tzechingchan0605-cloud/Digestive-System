"""Validate user-facing staged worksheets, not just their generation metadata."""
from pathlib import Path
from zipfile import ZipFile
import re,json,fitz
from docx import Document
root=Path(__file__).resolve().parents[1];out=root/'paperwork/generated'
rules=json.loads((root/'paperwork/rubric.json').read_text())
assert rules['worksheetDesignVersion']=='deductive-2'
assert rules['vlMaximum']==31 and rules['worksheetMaximum']==16
assert len(rules['reflectionPoints'])==8
for lang in ['','_en']:
 for stem in ['student','answers']:
  name=f'WS1_{stem}{lang}';doc=Document(out/(name+'.docx'));pdf=fitz.open(out/(name+'.pdf'))
  assert len(pdf)==4,(name,len(pdf))
  paragraphs=list(doc.paragraphs)
  for table in doc.tables:
   for row in table.rows:
    for cell in row.cells:paragraphs.extend(cell.paragraphs)
  text='\n'.join(p.text for p in paragraphs)
  assert len([p for p in doc.paragraphs if re.match(r'^\d+\. ',p.text)])==6
  assert '主要外觀' not in text and 'Appearance of X after' not in text
  assert '對照組' not in text and 'clear solution' not in text and '清澈' not in text
  assert len(doc.inline_shapes)==6
  reference=Document(root/f'paperwork/templates/WS1_topic_notes_{"en" if lang else "zh"}.docx')
  topic=next(i for i,p in enumerate(doc.paragraphs) if p.text in ['課題需知','Topic notes'])
  assert [p.text for p in doc.paragraphs[topic:topic+len(reference.paragraphs)]]==[p.text for p in reference.paragraphs]
  for actual,expected in zip(doc.paragraphs[topic:topic+len(reference.paragraphs)],reference.paragraphs):
   runs=lambda p:[(r.text,r.bold,str(r.font.color.rgb),r.font.size.pt if r.font.size else None) for r in p.runs]
   assert runs(actual)==runs(expected), 'Reference keyword formatting changed'
   if expected.style.font.bold:
    assert actual.style.font.bold==expected.style.font.bold
    assert str(actual.style.font.color.rgb)==str(expected.style.font.color.rgb)
  for actual,expected in zip(list(doc.inline_shapes)[:2],reference.inline_shapes):
   embed=lambda shape:shape._inline.xpath('.//a:blip')[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
   assert doc.part.related_parts[embed(actual)].blob==reference.part.related_parts[embed(expected)].blob
  assert doc.core_properties.subject.startswith('deductive-2')
  for p in paragraphs:
   for r in p.runs:
    if r.text.strip():assert (r.font.size or p.style.font.size or doc.styles['Normal'].font.size).pt>=11
  native=[''.join(s['text'] for s in l['spans']) for b in pdf[0].get_text('dict')['blocks'] if 'lines'in b for l in b['lines']]
  assert any(all(label in l for label in (['Name:','Class:','No.:'] if lang else ['姓名：','班別：','學號：'])) for l in native)
  # The deductive paper condition starts with concepts; the VL reveal stays at the end.
  for i,p in enumerate(pdf):
   t=p.get_text();assert len(t.strip())>100,(name,i,'orphan')
   if lang:assert not re.search(r'[\u3400-\u9fff]',t)
   assert abs(p.rect.width-595)<2 and abs(p.rect.height-842)<2
  first=pdf[0].get_text();assert all(w in first for w in (['bile','pancreatic','lipase'] if lang else ['膽汁','胰液','脂肪酶']))
  assert any(str(r.font.color.rgb)=='B7312C' and r.bold for p in doc.paragraphs for r in p.runs)
  if stem=='student':
   assert not any(p.text.startswith('☑') for p in doc.paragraphs)
   borders=[b for p in doc.paragraphs for b in p._p.xpath('./w:pPr/w:pBdr/w:bottom')]
   assert len(borders)==5
   rules=[item for drawing in pdf[-1].get_drawings() for item in drawing['items'] if item[0]=='l' and abs(item[1].y-item[2].y)<1 and abs(item[1].x-item[2].x)>450]
   assert len(rules)>=5, 'Printed response lines must not merge into a single border'
   for b in borders:assert b.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color')=='333333'
  else:
   assert sum(p.text.startswith('☑') for p in doc.paragraphs)==6 # Three premises + three analyses.
   assert any(str(r.font.color.rgb)=='145DA0' for p in doc.paragraphs for r in p.runs)
  print('PASS:',name,'4 A4 pages, 6 tasks, 6 diagrams, >=11pt, same-line identity, concept-first deductive flow, distinct key.')
 for stem in ['D_teacher_guide']:
  pdf=fitz.open(out/f'WS1_{stem}{lang}.pdf');t='\n'.join(p.get_text() for p in pdf)
  assert 'staged-1' in t and '31' in t and '16' in t
  if lang:assert not re.search(r'[\u3400-\u9fff]',t)
print('PASS: bilingual teacher guides, new maxima and dark response lines.')
